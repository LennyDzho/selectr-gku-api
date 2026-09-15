from __future__ import annotations

import math
import os
import re
import sys
import tempfile
from zipfile import BadZipFile
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Callable, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException


def _application_directory() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


APPLICATION_DIRECTORY = _application_directory()
INPUT_DIRECTORY = APPLICATION_DIRECTORY / "input"
OUTPUT_FILE = APPLICATION_DIRECTORY / "output" / "output.xlsx"

CREATED_AT = "Дата создания"
LOGIN = "Логин"
ROUTE = "Номер маршрута"
BOARD = "Бортовой номер"
TICKET_ID = "Уникальный номер билета"
TICKET_TYPE = "Название типа билета"

REQUIRED_HEADERS = (CREATED_AT, LOGIN, ROUTE, BOARD, TICKET_ID, TICKET_TYPE)
OUTPUT_HEADERS = (
    "Дата",
    "Время начала проверки ТС",
    LOGIN,
    ROUTE,
    BOARD,
    "Кол-во проверок",
    "Кол-во проверок ТС",
)

ZERO_TICKET_ID = "00000000000000000000"
SOCIAL_CARD_TYPE = "социальная карта москвича"
SESSION_GAP = timedelta(minutes=10)
DUPLICATE_GAP = timedelta(seconds=1)


class ReportError(Exception):
    """Base error for an expected report-generation failure."""


class WorkbookSchemaError(ReportError):
    """Raised when an input workbook cannot be mapped to the required schema."""


class NoValidDataError(ReportError):
    """Raised when no report can be produced from the available input."""


@dataclass(frozen=True, slots=True)
class Record:
    timestamp: datetime
    login: str
    route: str
    board: str
    ticket_id: str
    ticket_type: str = ""
    source_name: str = ""
    source_row: int = 0


@dataclass(frozen=True, slots=True)
class Visit:
    start: datetime
    login: str
    route: str
    board: str
    check_count: int
    visit_number: int
    visit_total: int

@dataclass(frozen=True, slots=True)
class ProcessingSummary:
    input_files: int
    processed_files: int
    skipped_files: int
    valid_rows: int
    skipped_rows: int
    visits: int
    checks: int


WarningHandler = Callable[[str], None]


def warn(message: str) -> None:
    print(f"Предупреждение: {message}", file=sys.stderr)


def _normalise_header(value: object) -> str:
    return "" if value is None else str(value).strip()


def _cell_to_text(cell: Cell) -> str:
    """Convert an identifier cell to text without stripping existing leading zeroes."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float):
        if not math.isfinite(value):
            return ""
        text = str(int(value)) if value.is_integer() else format(value, ".15g")
    else:
        return str(value).strip()

    # A numeric cell formatted as 00000000 is an Excel representation of a
    # text-like identifier. Restore the display zeroes when that is unambiguous.
    number_format = str(cell.number_format or "").split(";", 1)[0]
    if re.fullmatch(r"0+", number_format) and not text.startswith("-"):
        text = text.zfill(len(number_format))
    return text


def _parse_timestamp(cell: Cell, workbook_epoch: datetime) -> datetime:
    value = cell.value
    parsed: datetime

    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    elif isinstance(value, (int, float)) and not isinstance(value, bool) and cell.is_date:
        converted = from_excel(value, workbook_epoch)
        parsed = (
            converted
            if isinstance(converted, datetime)
            else datetime.combine(converted, time.min)
        )
    elif isinstance(value, str):
        text_value = value.strip()
        if not text_value:
            raise ValueError("пустая дата создания")

        normalised = text_value.replace(",", ".")
        try:
            parsed = datetime.fromisoformat(normalised)
        except ValueError:
            parsed = _parse_timestamp_with_known_formats(normalised)
    else:
        raise ValueError(f"неподдерживаемое значение даты: {value!r}")

    if parsed.tzinfo is not None:
        raise ValueError("дата создания с часовым поясом не поддерживается")
    return parsed


def _parse_timestamp_with_known_formats(value: str) -> datetime:
    formats = (
        "%d.%m.%Y %H:%M:%S.%f",
        "%d.%m.%Y %H:%M:%S",
        "%Y/%m/%d %H:%M:%S.%f",
        "%Y/%m/%d %H:%M:%S",
    )
    for timestamp_format in formats:
        try:
            return datetime.strptime(value, timestamp_format)
        except ValueError:
            continue
    raise ValueError(f"не удалось распознать дату создания {value!r}")


def read_records(
    workbook_path: Path,
    warning_handler: WarningHandler = warn,
) -> tuple[list[Record], int]:
    """Read valid records from the active sheet and return them with a skip count."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows()
        header_row = next(rows, None)
        if header_row is None:
            raise WorkbookSchemaError("активный лист пуст")

        header_positions: dict[str, int] = {}
        duplicate_headers: set[str] = set()
        for index, cell in enumerate(header_row):
            header = _normalise_header(cell.value)
            if header not in REQUIRED_HEADERS:
                continue
            if header in header_positions:
                duplicate_headers.add(header)
            else:
                header_positions[header] = index

        if duplicate_headers:
            duplicates = ", ".join(sorted(duplicate_headers))
            raise WorkbookSchemaError(f"повторяются обязательные колонки: {duplicates}")

        missing_headers = [
            header for header in REQUIRED_HEADERS if header not in header_positions
        ]
        if missing_headers:
            raise WorkbookSchemaError(
                "нет обязательных колонок: " + ", ".join(missing_headers)
            )

        records: list[Record] = []
        skipped_rows = 0
        for row_number, row in enumerate(rows, start=2):
            cells = {header: row[index] for header, index in header_positions.items()}

            # Styled but otherwise empty trailing rows are not data errors.
            if all(cell.value is None for cell in cells.values()):
                continue

            errors: list[str] = []
            try:
                timestamp = _parse_timestamp(cells[CREATED_AT], workbook.epoch)
            except (TypeError, ValueError) as error:
                timestamp = None
                errors.append(str(error))

            login = _cell_to_text(cells[LOGIN])
            route = _cell_to_text(cells[ROUTE])
            board = _cell_to_text(cells[BOARD])
            ticket_id = _cell_to_text(cells[TICKET_ID])
            ticket_type = _cell_to_text(cells[TICKET_TYPE])

            for field_name, field_value in (
                (LOGIN, login),
                (ROUTE, route),
                (BOARD, board),
                (TICKET_ID, ticket_id),
            ):
                if not field_value:
                    errors.append(f"пустое поле «{field_name}»")

            if errors:
                skipped_rows += 1
                warning_handler(
                    f"{workbook_path.name}, строка {row_number}: " + "; ".join(errors)
                )
                continue

            assert timestamp is not None
            records.append(
                Record(
                    timestamp=timestamp,
                    login=login,
                    route=route,
                    board=board,
                    ticket_id=ticket_id,
                    ticket_type=ticket_type,
                    source_name=workbook_path.name,
                    source_row=row_number,
                )
            )

        return records, skipped_rows
    finally:
        workbook.close()


def count_passenger_checks(records: Sequence[Record]) -> int:
    """Count passenger checks inside one vehicle visit."""
    if not records:
        return 0

    # A non-zero identifier is counted independently from every other
    # identifier. Repeated rows for that identifier are one transaction only
    # while consecutive timestamps differ by no more than one second.
    ticket_times: dict[str, list[datetime]] = defaultdict(list)
    social_card_seconds: set[datetime] = set()
    zero_ticket_seconds: list[datetime] = []

    for record in records:
        second = record.timestamp.replace(microsecond=0)
        if record.ticket_id == ZERO_TICKET_ID:
            zero_ticket_seconds.append(second)
            continue

        ticket_times[record.ticket_id].append(record.timestamp)
        if " ".join(record.ticket_type.split()).casefold() == SOCIAL_CARD_TYPE:
            social_card_seconds.add(second)

    # Zero identifiers normally represent separate transactions. They are
    # artifacts of a social-card scan only when a non-zero social-card record
    # exists on the same vehicle in the same second.
    check_count = sum(
        second not in social_card_seconds for second in zero_ticket_seconds
    )

    for timestamps in ticket_times.values():
        timestamps.sort()
        check_count += 1
        for previous, current in zip(timestamps, timestamps[1:]):
            if current - previous > DUPLICATE_GAP:
                check_count += 1

    return check_count


def build_visits(records: Iterable[Record]) -> list[Visit]:
    """Group records into daily inspector/route/board vehicle visits."""
    groups: dict[tuple[date, str, str, str], list[Record]] = defaultdict(list)
    for record in records:
        key = (record.timestamp.date(), record.login, record.route, record.board)
        groups[key].append(record)

    visits: list[Visit] = []
    for (_day, login, route, board), group_records in groups.items():
        ordered_records = sorted(
            group_records,
            key=lambda item: (
                item.timestamp,
                item.source_name.casefold(),
                item.source_row,
            ),
        )

        sessions: list[list[Record]] = []
        current_session: list[Record] = []
        for record in ordered_records:
            if (
                current_session
                and record.timestamp - current_session[-1].timestamp >= SESSION_GAP
            ):
                sessions.append(current_session)
                current_session = []
            current_session.append(record)
        if current_session:
            sessions.append(current_session)

        visit_total = len(sessions)
        for visit_number, session in enumerate(sessions, start=1):
            visits.append(
                Visit(
                    start=session[0].timestamp,
                    login=login,
                    route=route,
                    board=board,
                    check_count=count_passenger_checks(session),
                    visit_number=visit_number,
                    visit_total=visit_total,
                )
            )

    visits.sort(
        key=lambda visit: (
            visit.start,
            visit.login.casefold(),
            visit.route.casefold(),
            visit.board.casefold(),
            visit.visit_number,
        )
    )
    return visits


def write_report(visits: Sequence[Visit], output_path: Path) -> None:
    """Write a formatted workbook and atomically replace the previous report."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Проверки"
    worksheet.append(OUTPUT_HEADERS)

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for visit in visits:
        worksheet.append(
            (
                visit.start.date(),
                visit.start.time(),
                visit.login,
                visit.route,
                visit.board,
                visit.check_count,
                visit.visit_total,
            )
        )

    for row_number in range(2, worksheet.max_row + 1):
        worksheet.cell(row_number, 1).number_format = "yyyy-mm-dd"
        worksheet.cell(row_number, 2).number_format = "hh:mm:ss.000"
        worksheet.cell(row_number, 4).number_format = "@"
        worksheet.cell(row_number, 5).number_format = "@"

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:G{worksheet.max_row}"
    column_widths = (13, 27, 20, 18, 18, 20, 16)
    for column_cells, width in zip(worksheet.columns, column_widths):
        worksheet.column_dimensions[column_cells[0].column_letter].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=".output-",
            suffix=".xlsx",
            dir=output_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
        temporary_path = None
    finally:
        workbook.close()
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def generate_report(
    input_directory: Path = INPUT_DIRECTORY,
    output_file: Path = OUTPUT_FILE,
    warning_handler: WarningHandler = warn,
) -> ProcessingSummary:
    input_files = sorted(
        (
            path
            for path in input_directory.glob("*.xlsx")
            if path.is_file() and not path.name.startswith("~$")
        ),
        key=lambda path: path.name.casefold(),
    )

    if not input_files:
        raise NoValidDataError(f"в папке {input_directory} нет файлов .xlsx")

    all_records: list[Record] = []
    skipped_rows = 0
    processed_files = 0
    skipped_files = 0

    for workbook_path in input_files:
        try:
            records, file_skipped_rows = read_records(
                workbook_path, warning_handler=warning_handler
            )
        except (
            OSError,
            ValueError,
            BadZipFile,
            InvalidFileException,
            WorkbookSchemaError,
        ) as error:
            skipped_files += 1
            warning_handler(f"{workbook_path.name}: файл пропущен: {error}")
            continue

        processed_files += 1
        skipped_rows += file_skipped_rows
        all_records.extend(records)

    if not all_records:
        raise NoValidDataError("во входных файлах нет валидных строк")

    visits = build_visits(all_records)
    write_report(visits, output_file)

    return ProcessingSummary(
        input_files=len(input_files),
        processed_files=processed_files,
        skipped_files=skipped_files,
        valid_rows=len(all_records),
        skipped_rows=skipped_rows,
        visits=len(visits),
        checks=sum(visit.check_count for visit in visits),
    )


def main() -> int:
    try:
        summary = generate_report()
    except (ReportError, OSError) as error:
        print(f"Ошибка: {error}", file=sys.stderr)
        return 1

    print(f"Входных файлов: {summary.input_files}")
    print(f"Обработано файлов: {summary.processed_files}")
    print(f"Пропущено файлов: {summary.skipped_files}")
    print(f"Валидных строк: {summary.valid_rows}")
    print(f"Пропущено строк: {summary.skipped_rows}")
    print(f"Проверок ТС: {summary.visits}")
    print(f"Проверок пассажиров: {summary.checks}")
    print(f"Отчёт сохранён: {OUTPUT_FILE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
