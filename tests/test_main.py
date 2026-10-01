from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from main import (
    NoValidDataError,
    OUTPUT_HEADERS,
    VISIT_START,
    ZERO_TICKET_ID,
    Record,
    WorkbookSchemaError,
    build_daily_report_rows,
    build_visits,
    count_passenger_checks,
    generate_report,
    read_records,
)


INPUT_HEADERS = (
    "Дата создания",
    "Логин",
    "Номер маршрута",
    "Бортовой номер",
    "Уникальный номер билета",
    "Название типа билета",
    VISIT_START,
)


def record(
    timestamp: str,
    ticket_id: str,
    board: str = "00193554",
    ticket_type: str = "",
    *,
    visit_start: str | None = None,
    login: str = "biryukovlb",
    route: str = "1542",
) -> Record:
    parsed_timestamp = datetime.fromisoformat(timestamp)
    return Record(
        timestamp=parsed_timestamp,
        login=login,
        route=route,
        board=board,
        ticket_id=ticket_id,
        visit_start=(
            datetime.fromisoformat(visit_start)
            if visit_start is not None
            else parsed_timestamp.replace(hour=0, minute=0, second=0, microsecond=0)
        ),
        ticket_type=ticket_type,
    )


class CheckCountingTests(unittest.TestCase):
    def test_social_card_absorbs_only_zero_rows_from_the_same_second(self) -> None:
        records = [
            record("2026-08-26 12:59:33", ZERO_TICKET_ID),
            record("2026-08-26 12:59:33", ZERO_TICKET_ID),
            record(
                "2026-08-26 12:59:33",
                "34139E11412806",
                ticket_type="Социальная карта москвича",
            ),
            record(
                "2026-08-26 12:59:33",
                "34139E11412806",
                ticket_type="Социальная карта москвича",
            ),
            record(
                "2026-08-26 12:59:33",
                "34139E11412803",
                ticket_type="Тройка",
            ),
        ]

        self.assertEqual(count_passenger_checks(records), 2)

    def test_social_card_does_not_absorb_another_nonzero_ticket(self) -> None:
        records = [
            record(
                "2026-08-26 12:59:33",
                "34139E11412806",
                ticket_type="Социальная карта москвича",
            ),
            record(
                "2026-08-26 12:59:33",
                "34139E11412803",
                ticket_type="Тройка",
            ),
        ]

        self.assertEqual(count_passenger_checks(records), 2)

    def test_nonzero_ticket_is_counted_once_despite_large_time_gaps(self) -> None:
        records = [
            record("2026-08-26 08:00:00.000", "TICKET"),
            record("2026-08-26 08:00:01.000", "TICKET"),
            record("2026-08-26 08:00:02.001", "TICKET"),
            record("2026-08-26 09:00:00", "TICKET"),
        ]

        self.assertEqual(count_passenger_checks(records), 1)

    def test_each_zero_ticket_row_is_counted(self) -> None:
        records = [
            record("2026-08-26 08:00:00", ZERO_TICKET_ID),
            record("2026-08-26 08:00:00", ZERO_TICKET_ID),
            record("2026-08-26 08:00:01", ZERO_TICKET_ID),
        ]

        self.assertEqual(count_passenger_checks(records), 3)

    def test_social_card_uses_calendar_seconds_not_elapsed_seconds(self) -> None:
        for zero_time, social_time, expected in (
            ("12:59:33.100", "12:59:33.900", 1),
            ("12:59:33.900", "12:59:34.000", 2),
        ):
            with self.subTest(zero_time=zero_time, social_time=social_time):
                records = [
                    record(f"2026-08-26 {zero_time}", ZERO_TICKET_ID),
                    record(
                        f"2026-08-26 {social_time}",
                        "SOCIAL",
                        ticket_type="Социальная карта москвича",
                    ),
                ]
                self.assertEqual(count_passenger_checks(records), expected)

    def test_social_card_type_requires_exact_normalised_name(self) -> None:
        for ticket_type, expected in (
            (" СОЦИАЛЬНАЯ   КАРТА МОСКВИЧА ", 1),
            ("Социальная карта москвича ММ", 2),
            ("", 2),
        ):
            with self.subTest(ticket_type=ticket_type):
                records = [
                    record("2026-08-26 12:59:33", ZERO_TICKET_ID),
                    record(
                        "2026-08-26 12:59:33", "SOCIAL", ticket_type=ticket_type
                    ),
                ]
                self.assertEqual(count_passenger_checks(records), expected)

    def test_zero_ticket_cannot_trigger_the_social_card_exception(self) -> None:
        records = [
            record("2026-08-26 12:59:33", ZERO_TICKET_ID),
            record(
                "2026-08-26 12:59:33",
                ZERO_TICKET_ID,
                ticket_type="Социальная карта москвича",
            ),
        ]
        self.assertEqual(count_passenger_checks(records), 2)


class VisitGroupingTests(unittest.TestCase):
    def test_same_ticket_and_second_on_different_vehicles_are_separate(self) -> None:
        visits = build_visits(
            [
                record("2026-08-26 08:00:00", "A", board="00193554"),
                record("2026-08-26 08:00:00", "A", board="00193555"),
            ]
        )

        self.assertEqual(len(visits), 2)
        self.assertEqual(sum(visit.check_count for visit in visits), 2)

    def test_same_source_start_keeps_one_visit_despite_large_pauses(self) -> None:
        visits = build_visits(
            [
                record(
                    "2026-08-26 08:00:00", "A", visit_start="2026-08-26 07:55:00"
                ),
                record(
                    "2026-08-26 08:09:59", "A", visit_start="2026-08-26 07:55:00"
                ),
                record(
                    "2026-08-26 09:00:00", "A", visit_start="2026-08-26 07:55:00"
                ),
            ]
        )

        self.assertEqual(len(visits), 1)
        self.assertEqual(visits[0].start, datetime(2026, 8, 26, 7, 55))
        self.assertEqual(visits[0].check_count, 1)

    def test_one_source_visit_can_cross_midnight(self) -> None:
        visits = build_visits(
            [
                record(
                    "2026-08-26 23:59:59", "A", visit_start="2026-08-26 23:58:00"
                ),
                record(
                    "2026-08-27 00:00:00", "A", visit_start="2026-08-26 23:58:00"
                ),
            ]
        )

        self.assertEqual(len(visits), 1)
        self.assertEqual(visits[0].check_count, 1)
        report_rows = build_daily_report_rows(visits)
        self.assertEqual(len(report_rows), 1)
        self.assertEqual(report_rows[0].day, datetime(2026, 8, 26).date())
        self.assertEqual(report_rows[0].vehicle_check_count, 1)

    def test_different_starts_create_visits_even_one_second_apart(self) -> None:
        visits = build_visits(
            [
                record(
                    "2026-08-26 08:00:00", "A", visit_start="2026-08-26 07:59:00"
                ),
                record(
                    "2026-08-26 08:00:01", "A", visit_start="2026-08-26 08:00:00"
                ),
            ]
        )
        self.assertEqual(len(visits), 2)
        self.assertEqual([visit.check_count for visit in visits], [1, 1])
        self.assertEqual([visit.visit_total for visit in visits], [2, 2])

    def test_source_start_milliseconds_distinguish_visits(self) -> None:
        visits = build_visits(
            [
                record(
                    "2026-08-26 08:00:01",
                    "A",
                    visit_start="2026-08-26 08:00:00.100",
                ),
                record(
                    "2026-08-26 08:00:01",
                    "A",
                    visit_start="2026-08-26 08:00:00.101",
                ),
            ]
        )
        self.assertEqual(len(visits), 2)
        self.assertEqual([visit.start.microsecond for visit in visits], [100000, 101000])
        self.assertEqual(sum(visit.check_count for visit in visits), 2)

    def test_source_start_is_also_scoped_to_login_route_and_board(self) -> None:
        visits = build_visits(
            [
                record("2026-08-26 08:00:00", "A"),
                record("2026-08-26 08:00:00", "A", login="other-inspector"),
                record("2026-08-26 08:00:00", "A", route="other-route"),
                record("2026-08-26 08:00:00", "A", board="other-board"),
            ]
        )
        self.assertEqual(len(visits), 4)
        self.assertEqual(sum(visit.check_count for visit in visits), 4)

    def test_social_card_cannot_absorb_zero_rows_from_another_visit(self) -> None:
        visits = build_visits(
            [
                record(
                    "2026-08-26 08:01:00",
                    ZERO_TICKET_ID,
                    visit_start="2026-08-26 07:55:00",
                ),
                record(
                    "2026-08-26 08:01:00",
                    "SOCIAL",
                    ticket_type="Социальная карта москвича",
                    visit_start="2026-08-26 08:00:00",
                ),
            ]
        )
        self.assertEqual(len(visits), 2)
        self.assertEqual([visit.check_count for visit in visits], [1, 1])

    def test_daily_report_combines_all_visits_of_the_same_vehicle(self) -> None:
        visits = build_visits(
            [
                record("2026-08-26 08:00:00", "A"),
                record("2026-08-26 08:01:00", "B"),
                record(
                    "2026-08-26 08:11:00", "C", visit_start="2026-08-26 08:10:00"
                ),
            ]
        )

        report_rows = build_daily_report_rows(visits)

        self.assertEqual(len(report_rows), 1)
        self.assertEqual(report_rows[0].check_count, 3)
        self.assertEqual(report_rows[0].vehicle_check_count, 2)

    def test_daily_rows_sort_by_source_day_login_route_and_board(self) -> None:
        visits = build_visits(
            [
                record("2026-08-27 08:00:00", "A", login="a", route="1", board="01"),
                record("2026-08-26 08:00:00", "A", login="b", route="1", board="01"),
                record("2026-08-26 08:00:00", "A", login="a", route="2", board="01"),
                record("2026-08-26 08:00:00", "A", login="a", route="1", board="02"),
                record("2026-08-26 08:00:00", "A", login="a", route="1", board="01"),
            ]
        )
        report_rows = build_daily_report_rows(visits)
        self.assertEqual(
            [(row.day.isoformat(), row.login, row.route, row.board) for row in report_rows],
            [
                ("2026-08-26", "a", "1", "01"),
                ("2026-08-26", "a", "1", "02"),
                ("2026-08-26", "a", "2", "01"),
                ("2026-08-26", "b", "1", "01"),
                ("2026-08-27", "a", "1", "01"),
            ],
        )


class ReportIntegrationTests(unittest.TestCase):
    def test_multiple_files_produce_one_sorted_report_and_skip_bad_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            input_directory = root / "input"
            output_file = root / "output" / "output.xlsx"
            input_directory.mkdir()

            self._write_input(
                input_directory / "later.xlsx",
                [
                    (
                        "2026-08-27 10:00:00",
                        "user2",
                        "2",
                        "00000002",
                        "B",
                        "Undefined type",
                        "2026-08-27 09:55:00",
                    ),
                    (
                        "not-a-date",
                        "user2",
                        "2",
                        "00000002",
                        "C",
                        "Undefined type",
                        "2026-08-27 09:55:00",
                    ),
                ],
            )
            self._write_input(
                input_directory / "earlier.xlsx",
                [
                    (
                        "2026-08-26 09:00:00",
                        "user1",
                        "1",
                        "00000001",
                        "A",
                        "Undefined type",
                        "2026-08-26 08:55:00",
                    )
                ],
            )

            warnings: list[str] = []
            summary = generate_report(input_directory, output_file, warnings.append)

            self.assertEqual(summary.input_files, 2)
            self.assertEqual(summary.valid_rows, 2)
            self.assertEqual(summary.skipped_rows, 1)
            self.assertEqual(summary.visits, 2)
            self.assertEqual(summary.checks, 2)
            self.assertEqual(len(warnings), 1)
            self.assertIn("Дата создания", warnings[0])
            self.assertTrue(output_file.exists())
            self.assertEqual(list(output_file.parent.glob("*.xlsx")), [output_file])

            workbook = load_workbook(output_file, data_only=True)
            try:
                worksheet = workbook["Проверки"]
                self.assertEqual(
                    tuple(cell.value for cell in worksheet[1]), OUTPUT_HEADERS
                )
                self.assertEqual(worksheet["B2"].value, "user1")
                self.assertEqual(worksheet["B3"].value, "user2")
                self.assertEqual(worksheet["C2"].value, "1")
                self.assertEqual(worksheet["D2"].value, "00000001")
                self.assertEqual(worksheet["E2"].value, 1)
                self.assertEqual(worksheet["F2"].value, 1)
            finally:
                workbook.close()

    def test_multiple_files_merge_source_visits_and_keep_one_daily_row(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            input_directory = root / "input"
            output_file = root / "output" / "output.xlsx"
            input_directory.mkdir()
            self._write_input(
                input_directory / "first.xlsx",
                [
                    (
                        "2026-08-26 08:00:00",
                        "user",
                        "1542",
                        "00193554",
                        "A",
                        "Тройка",
                        "2026-08-26 07:55:00",
                    ),
                ],
            )
            self._write_input(
                input_directory / "second.xlsx",
                [
                    (
                        "2026-08-26 08:20:00",
                        "user",
                        "1542",
                        "00193554",
                        "A",
                        "Кошелек",
                        "2026-08-26 07:55:00",
                    ),
                    (
                        "2026-08-26 08:22:00",
                        "user",
                        "1542",
                        "00193554",
                        "A",
                        "Тройка",
                        "2026-08-26 08:21:00",
                    ),
                ],
            )
            summary = generate_report(input_directory, output_file)
            self.assertEqual(summary.input_files, 2)
            self.assertEqual(summary.valid_rows, 3)
            self.assertEqual(summary.visits, 2)
            self.assertEqual(summary.checks, 2)
            workbook = load_workbook(output_file, data_only=True)
            try:
                worksheet = workbook["Проверки"]
                self.assertEqual(worksheet.max_row, 2)
                self.assertEqual(worksheet["E2"].value, 2)
                self.assertEqual(worksheet["F2"].value, 2)
            finally:
                workbook.close()

    def test_missing_or_duplicate_source_start_header_is_rejected(self) -> None:
        for headers in (INPUT_HEADERS[:-1], INPUT_HEADERS + (VISIT_START,)):
            with self.subTest(headers=headers):
                with tempfile.TemporaryDirectory() as temporary_directory:
                    path = Path(temporary_directory) / "bad-header.xlsx"
                    self._write_input(path, [], headers=headers)
                    with self.assertRaisesRegex(WorkbookSchemaError, VISIT_START):
                        read_records(path)

    def test_invalid_source_start_headers_are_skipped_without_losing_valid_files(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            input_directory = root / "input"
            output_file = root / "output" / "output.xlsx"
            input_directory.mkdir()
            valid_row = (
                "2026-08-26 08:00:00",
                "user",
                "1542",
                "00193554",
                "A",
                "Тройка",
                "2026-08-26 07:55:00",
            )
            self._write_input(input_directory / "valid.xlsx", [valid_row])
            self._write_input(input_directory / "~$temporary.xlsx", [valid_row])
            self._write_input(
                input_directory / "missing-start.xlsx", [], headers=INPUT_HEADERS[:-1]
            )
            self._write_input(
                input_directory / "duplicate-start.xlsx",
                [],
                headers=INPUT_HEADERS + (VISIT_START,),
            )
            warnings: list[str] = []
            summary = generate_report(input_directory, output_file, warnings.append)
            self.assertEqual(summary.input_files, 3)
            self.assertEqual(summary.processed_files, 1)
            self.assertEqual(summary.skipped_files, 2)
            self.assertEqual(summary.checks, 1)
            self.assertEqual(len(warnings), 2)
            self.assertTrue(all(VISIT_START in warning for warning in warnings))
            self.assertTrue(output_file.exists())

    def test_invalid_or_blank_source_start_is_skipped_with_field_and_row(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "invalid-start.xlsx"
            self._write_input(
                path,
                [
                    (
                        "2026-08-26 08:00:00",
                        "user",
                        "1542",
                        "00193554",
                        "A",
                        "Тройка",
                        visit_start,
                    )
                    for visit_start in ("not-a-date", None, " ")
                ],
            )
            warnings: list[str] = []
            records, skipped = read_records(path, warnings.append)
            self.assertEqual(records, [])
            self.assertEqual(skipped, 3)
            self.assertEqual(len(warnings), 3)
            for row_number, warning in enumerate(warnings, start=2):
                self.assertIn(path.name, warning)
                self.assertIn(f"строка {row_number}", warning)
                self.assertIn(VISIT_START, warning)

    def test_both_timestamps_can_be_excel_dates_and_keep_milliseconds(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "excel-dates.xlsx"
            timestamp = datetime(2026, 8, 26, 8, 0, 0, 456000)
            visit_start = datetime(2026, 8, 26, 7, 55, 0, 123000)
            self._write_input(
                path,
                [(timestamp, "user", "1542", "00193554", "A", "", visit_start)],
            )
            records, skipped = read_records(path)
            self.assertEqual(skipped, 0)
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0].timestamp, timestamp)
            self.assertEqual(records[0].visit_start, visit_start)

    def test_no_valid_rows_do_not_replace_an_existing_report(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            input_directory = root / "input"
            output_file = root / "output" / "output.xlsx"
            input_directory.mkdir()
            output_file.parent.mkdir()
            output_file.write_bytes(b"existing report")
            self._write_input(
                input_directory / "invalid.xlsx",
                [
                    (
                        "not-a-date",
                        "user",
                        "1",
                        "00000001",
                        "A",
                        "Undefined type",
                        "2026-08-26 07:55:00",
                    )
                ],
            )

            with self.assertRaises(NoValidDataError):
                generate_report(input_directory, output_file, lambda _message: None)

            self.assertEqual(output_file.read_bytes(), b"existing report")

    def test_no_valid_source_starts_do_not_replace_an_existing_report(self) -> None:
        for invalid_start in (None, "not-a-date"):
            with self.subTest(invalid_start=invalid_start):
                with tempfile.TemporaryDirectory() as temporary_directory:
                    root = Path(temporary_directory)
                    input_directory = root / "input"
                    output_file = root / "output" / "output.xlsx"
                    input_directory.mkdir()
                    output_file.parent.mkdir()
                    output_file.write_bytes(b"existing report")
                    self._write_input(
                        input_directory / "invalid.xlsx",
                        [
                            (
                                "2026-08-26 08:00:00",
                                "user",
                                "1542",
                                "00193554",
                                "A",
                                "Тройка",
                                invalid_start,
                            )
                        ],
                    )
                    with self.assertRaises(NoValidDataError):
                        generate_report(input_directory, output_file, lambda _: None)
                    self.assertEqual(output_file.read_bytes(), b"existing report")

    @staticmethod
    def _write_input(
        path: Path,
        rows: list[tuple[object, ...]],
        *,
        headers: tuple[str, ...] = INPUT_HEADERS,
    ) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
